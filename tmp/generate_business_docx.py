from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pathlib import Path

OUT = Path('artifacts/business_deviation_table_refined.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

rows = [
('资格性审查','独立承担民事责任的能力','比选申请人应具有独立承担民事责任的能力。','比选申请人为依法设立并有效存续的法人单位，具有独立承担民事责任的能力，完全满足要求。'),
('资格性审查','商业信誉和财务会计制度','比选申请人应具有良好的商业信誉和健全的财务会计制度。','比选申请人具有良好的商业信誉和健全的财务会计制度，完全满足要求。'),
('资格性审查','履行合同所需设备和专业技术能力','比选申请人应具有履行合同所必需的设备和专业技术能力。','比选申请人具备完成本项目所需的科研条件、软硬件设备、专业技术团队和项目实施能力，完全满足要求。'),
('资格性审查','依法缴纳税收','比选申请人应具有依法缴纳税收的良好记录。','比选申请人依法履行纳税义务，具有良好的税收缴纳记录，完全满足要求。'),
('资格性审查','依法缴纳社会保障资金','比选申请人应具有依法缴纳社会保障资金的良好记录。','比选申请人依法缴纳社会保障资金，具有良好的社会保障资金缴纳记录，完全满足要求。'),
('资格性审查','无重大违法记录','比选申请人参加政府采购活动前三年内，在经营活动中没有重大违法记录。','比选申请人参加政府采购活动前三年内，在经营活动中没有重大违法记录，并按要求提供书面声明。'),
('资格性审查','法律、行政法规规定的其他条件','比选申请人应符合法律、行政法规规定的其他条件。','比选申请人符合法律、行政法规规定的其他资格条件，完全满足要求。'),
('资格性审查','法人证明文件','提供有效的营业执照复印件并加盖公章；事业单位应提供有效的事业单位法人证书复印件。','比选申请人将提供有效的法人证明文件复印件并加盖单位公章，完全满足要求。'),
('资格性审查','良好商业信誉及信用记录','提供“信用中国”网站中失信被执行人查询截图及承诺书；比选申请人不得被列入相关失信或政府采购严重违法失信名单。','比选申请人未被列入相关失信名单，将按要求提供信用查询截图及相关承诺材料。'),
('资格性审查','其他法定资格条件','不存在法律、行政法规规定的其他不得参与本项目比选的情形。','比选申请人不存在法律、行政法规规定的其他不得参与本项目比选的情形。'),
('符合性审查','申请文件格式','按照比选文件规定的格式和内容填写，字迹清晰可辨，资料完整。','比选申请文件按照比选文件规定的格式和内容编制，文字清晰，资料完整。'),
('符合性审查','申请文件签字盖章','比选申请文件应有法定代表人或法定代表人授权代理人签字，并加盖比选申请人公章。','比选申请文件将按照要求完成法定代表人或授权代理人签字及单位盖章。'),
('符合性审查','报价唯一','每个比选申请人只能有一个有效报价，不得提供选择性报价。','比选申请人仅提供一个明确、唯一且有效的报价，不提供选择性报价。'),
('符合性审查','比选有效期','比选有效期为自比选申请文件递交截止之日起90日历天。','比选申请文件有效期为自递交截止之日起90日历天，完全满足要求。'),
('符合性审查','项目周期','自合同签订之日起5个月内完成全部技术服务。','比选申请人承诺自合同签订之日起5个月内完成全部技术服务及成果验收工作。'),
('符合性审查','其他无效情形','比选申请人及比选申请文件不存在不符合法律、法规和比选文件规定的其他无效情形。','比选申请人及比选申请文件不存在法律、法规和比选文件规定的其他无效情形。'),
('商务评分','比选申请人研究实力','比选申请人应具有较强的研究实力，在时空知识图谱、国土空间格局分析、城市群空间绩效评估等领域具有相关研究经验。','比选申请人长期从事时空知识图谱、国土空间格局分析、城市空间绩效评估等相关研究，具备较强的科研基础和项目实施能力。'),
('商务评分','相关项目业绩一','近五年承担过时空知识图谱、国土空间格局分析、城市更新或城市群空间绩效评估等相关研究项目，并提供任务书或项目合同书。每项得2分，最高10分。','比选申请人拟提供第1项符合评分范围的相关研究项目及任务书或合同证明材料。'),
('商务评分','相关项目业绩二','近五年承担过时空知识图谱、国土空间格局分析、城市更新或城市群空间绩效评估等相关研究项目，并提供任务书或项目合同书。每项得2分，最高10分。','比选申请人拟提供第2项符合评分范围的相关研究项目及任务书或合同证明材料。'),
('商务评分','相关项目业绩三','近五年承担过时空知识图谱、国土空间格局分析、城市更新或城市群空间绩效评估等相关研究项目，并提供任务书或项目合同书。每项得2分，最高10分。','比选申请人拟提供第3项符合评分范围的相关研究项目及任务书或合同证明材料。'),
('商务评分','相关项目业绩四','近五年承担过时空知识图谱、国土空间格局分析、城市更新或城市群空间绩效评估等相关研究项目，并提供任务书或项目合同书。每项得2分，最高10分。','比选申请人拟提供第4项符合评分范围的相关研究项目及任务书或合同证明材料。'),
('商务评分','相关项目业绩五','近五年承担过时空知识图谱、国土空间格局分析、城市更新或城市群空间绩效评估等相关研究项目，并提供任务书或项目合同书。每项得2分，最高10分。','比选申请人拟提供第5项符合评分范围的相关研究项目及任务书或合同证明材料。'),
('商务评分','项目负责人论文成果','项目负责人以第一作者或通讯作者发表地理学、人口、交通、国土空间格局分析、遥感监测、信息抽取、知识图谱等相关领域高质量论文。发表8篇以上得10分，3—7篇得5分，2篇及以下不得分。','项目负责人已在相关领域发表符合评分要求的高质量论文，拟提供论文清单、论文首页及检索证明等支撑材料。'),
('商务评分','项目组专业背景','项目组成员应具备国土空间格局优化、GIS、地理学、知识图谱等方面的工作或学习经历，人员专业结构合理。','项目组成员专业背景涵盖地理学、地理信息、遥感科学、计算机科学、人工智能及知识图谱等方向，人员结构与项目任务高度匹配。'),
('商务评分','博士学位人员数量','项目组成员中2人具备地理学、地理信息、遥感科学、计算机科学、人工智能等相关学科博士学位的，得5分；须提供学位证书复印件。','项目组中不少于2名成员具备相关学科博士学位，并将提供学位证书复印件作为证明。'),
('商务评分','项目负责人职称','项目负责人具有副高级及以上职称的得5分，中级职称得1分，初级职称不得分；须提供职称证书复印件。','项目负责人具备副高级及以上职称，并将提供有效职称证书复印件作为证明。'),
('商务评分','商务材料完整性','项目业绩、论文、人员学位及项目负责人职称等材料应真实、完整，并提供相应证明文件。','比选申请人将按照评分标准逐项提供任务书或合同、论文证明、学位证书及职称证书等材料，保证材料真实、完整、有效。'),
]

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)

def set_font(run, size=9, bold=False, name='宋体'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

style = doc.styles['Normal']
style.font.name = '宋体'
style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.font.size = Pt(9)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('商务条款偏离表')
set_font(r, size=18, bold=True, name='方正小标宋简体')
p.paragraph_format.space_after = Pt(8)

for label, value in [('项目名称：','面向国土空间规划的语义关联网络构建与文本智能分析模型研究'),('项目编号：','HXLDZB-FW-20260226')]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    set_font(r, size=10.5, bold=True)
    r = p.add_run(value)
    set_font(r, size=10.5)

headers = ['序号','条款类别','评审因素','比选文件要求','比选申请人响应情况','响应文件位置','偏离情况']
table = doc.add_table(rows=1, cols=len(headers))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
table.autofit = False
widths = [Cm(1.0),Cm(2.0),Cm(2.8),Cm(7.6),Cm(8.0),Cm(3.1),Cm(2.0)]
for i,w in enumerate(widths):
    table.columns[i].width = w
hdr = table.rows[0]
set_repeat_table_header(hdr)
for i,text in enumerate(headers):
    cell = hdr.cells[i]
    cell.width = widths[i]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, 'E7E6E6')
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, size=9, bold=True)

for idx,(cat,factor,req,resp) in enumerate(rows, start=1):
    cells = table.add_row().cells
    values = [str(idx),cat,factor,req,resp,'【待补充】','无偏离']
    for i,(cell,text) in enumerate(zip(cells,values)):
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0,1,5,6) else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        set_font(r, size=8.5)

notes = [
'本表所列各项均为无偏离，且比选申请人对相关要求作完全响应。',
'“响应文件位置”应在比选申请文件完成编排和页码确定后，补充对应章节名称及页码。',
'第18—22项应分别对应5项可计分项目业绩，不宜合并填写，以便评审专家逐项核验和计分。',
'第23项建议在响应文件位置中同时注明“论文汇总表、论文首页及检索证明”的具体页码。',
'第25项和第26项应分别指向博士学位证书及项目负责人职称证书所在页码。',
]
p = doc.add_paragraph()
r = p.add_run('说明：')
set_font(r, size=10.5, bold=True)
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
for i,n in enumerate(notes, start=1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f'{i}. {n}')
    set_font(r, size=10)

for line in [
'比选申请人：人工智能与数字经济广东省实验室（深圳）（盖章）',
'法定代表人或其授权代理人：________________（签字或盖章）',
'日    期：________年____月____日']:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(line)
    set_font(r, size=10.5)

doc.save(OUT)
print(OUT)
