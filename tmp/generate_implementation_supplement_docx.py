from pathlib import Path
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path('tmp/implementation_plan_supplement_text_analysis_20260726.doc')
OUT = Path('artifacts/implementation_plan_supplement_text_analysis_20260726.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)


def set_run_font(run, east='宋体', latin='Times New Roman', size=12, bold=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def add_inline(paragraph, node, inherited_bold=False):
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, size=12, bold=inherited_bold)
        return
    if not isinstance(node, Tag):
        return
    if node.name == 'br':
        paragraph.add_run().add_break()
        return
    bold = inherited_bold or node.name in ('b', 'strong')
    for child in node.children:
        add_inline(paragraph, child, bold)


def style_paragraph(paragraph, *, align=None, indent=True, spacing=1.5, before=0, after=6):
    paragraph.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.line_spacing = spacing
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.first_line_indent = Cm(0.74) if indent else None


def add_heading_from_tag(doc, tag, level):
    p = doc.add_paragraph()
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        for child in tag.children:
            if isinstance(child, Tag) and child.name == 'br':
                p.add_run().add_break()
            else:
                add_inline(p, child, True)
        for run in p.runs:
            set_run_font(run, east='方正小标宋简体', size=20, bold=True)
        return
    p.style = f'Heading {level}'
    p.paragraph_format.keep_with_next = True
    for child in tag.children:
        add_inline(p, child, True)
    font_map = {1: ('黑体', 16), 2: ('黑体', 14), 3: ('楷体', 12)}
    east, size = font_map[level]
    for run in p.runs:
        set_run_font(run, east=east, size=size, bold=True)


def add_table(doc, tag):
    rows = tag.find_all('tr', recursive=True)
    if not rows:
        return
    max_cols = max(len(r.find_all(['th', 'td'], recursive=False)) for r in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, r in enumerate(rows):
        cells = r.find_all(['th', 'td'], recursive=False)
        for j in range(max_cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = cells[j].get_text(' ', strip=True) if j < len(cells) else ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or 'center' in (cells[j].get('class') or []) else WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_run_font(run, size=10.5, bold=(i == 0 or cells[j].name == 'th'))
            if i == 0 or cells[j].name == 'th':
                set_cell_shading(cell, 'E7E6E6')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def ensure_styles(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Times New Roman'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.font.size = Pt(12)
    for name, east, size in [('Heading 1', '黑体', 16), ('Heading 2', '黑体', 14), ('Heading 3', '楷体', 12)]:
        st = styles[name]
        st.font.name = 'Times New Roman'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), east)
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)


def main():
    soup = BeautifulSoup(SRC.read_text(encoding='utf-8'), 'html.parser')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2)
    sec.right_margin = Cm(2)
    ensure_styles(doc)

    body = soup.body
    for node in body.descendants:
        # Only process direct children of body or of the main Section1 div.
        if not isinstance(node, Tag):
            continue
        parent = node.parent
        if parent is body or (isinstance(parent, Tag) and parent.name == 'div' and 'Section1' in (parent.get('class') or [])):
            if node.name == 'h1':
                add_heading_from_tag(doc, node, 0)
            elif node.name == 'h2':
                add_heading_from_tag(doc, node, 1)
            elif node.name == 'h3':
                add_heading_from_tag(doc, node, 2)
            elif node.name == 'h4':
                add_heading_from_tag(doc, node, 3)
            elif node.name == 'p':
                classes = node.get('class') or []
                p = doc.add_paragraph()
                is_center = 'center' in classes
                is_note = 'note' in classes
                is_noindent = 'noindent' in classes or is_center or is_note
                style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER if is_center else WD_ALIGN_PARAGRAPH.JUSTIFY, indent=not is_noindent, spacing=1.5, after=6)
                for child in node.children:
                    add_inline(p, child)
                if is_note:
                    p.paragraph_format.left_indent = Cm(0.3)
                    p.paragraph_format.right_indent = Cm(0.3)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(8)
                    p_pr = p._p.get_or_add_pPr()
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:fill'), 'F2F2F2')
                    p_pr.append(shd)
                    pbdr = OxmlElement('w:pBdr')
                    for edge in ('top', 'left', 'bottom', 'right'):
                        el = OxmlElement(f'w:{edge}')
                        el.set(qn('w:val'), 'single')
                        el.set(qn('w:sz'), '4')
                        el.set(qn('w:color'), '999999')
                        pbdr.append(el)
                    p_pr.append(pbdr)
            elif node.name == 'table':
                add_table(doc, node)
            elif node.name == 'div' and 'pagebreak' in (node.get('class') or []):
                doc.add_page_break()

    # Footer with centered page number field.
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld_char1 = OxmlElement('w:fldChar')
        fld_char1.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = ' PAGE '
        fld_char2 = OxmlElement('w:fldChar')
        fld_char2.set(qn('w:fldCharType'), 'end')
        run._r.extend([fld_char1, instr, fld_char2])
        set_run_font(run, size=10.5)

    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
